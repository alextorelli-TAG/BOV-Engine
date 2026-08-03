#!/usr/bin/env python3
"""Excel ingestion for the BOV engine — comps tables and property proformas.

Server-side (openpyxl). The console uploads a workbook; these parsers turn the
real-world sheet shapes into the structured data the builders consume:

  parse_comps(data, kind)   -> {'rows': [ {name,address,city,price,bldg,ppu,cap,
                                            close,units,asset_type,notes}, ... ]}
  parse_financials(data)    -> {'property': {...console fields...},
                                'highlights': [...],
                                'operating': [ {label,amount,psf,group,total} ],
                                'cashflow': {'years': [...], 'rows': [...]} }

Both tolerate the messy reality of broker workbooks: title rows above the header,
a header that isn't on row 1, footer/source rows, and labelled key/value grids.
"""
import datetime
import io
import re

import openpyxl


# --------------------------------------------------------------- value helpers
def num(v):
    """Best-effort number from a cell that may be int/float or a string like
    '140,852 RSF' or '$1,207'. Returns float or None."""
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        m = re.search(r'-?\d[\d,]*\.?\d*', v.replace('$', ''))
        if m:
            try:
                return float(m.group(0).replace(',', ''))
            except ValueError:
                return None
    return None


def money(v):
    n = num(v)
    return f'${round(n):,}' if n is not None else ''


def money2(v):
    n = num(v)
    return f'${n:,.2f}' if n is not None else ''


def as_pct(v):
    n = num(v)
    if n is None:
        return ''
    if abs(n) <= 1:
        n *= 100
    return f'{n:.1f}%'


def date_str(v):
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime('%b %Y')
    s = str(v).strip() if v is not None else ''
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', s)
    if m:
        try:
            return datetime.date(int(m[1]), int(m[2]), int(m[3])).strftime('%b %Y')
        except ValueError:
            pass
    return s


def year_of(v):
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.year
    n = num(v)
    return int(n) if n is not None else None


def norm(s):
    return re.sub(r'[^a-z0-9]', '', str(s).lower()) if s is not None else ''


def _wb(data):
    """Load a workbook from bytes / file-like / path, computed values only."""
    if isinstance(data, (bytes, bytearray)):
        data = io.BytesIO(data)
    return openpyxl.load_workbook(data, data_only=True, read_only=True)


# ------------------------------------------------------------------ comps
# header synonym -> canonical key (normalised, so "Property Address" -> address)
COMP_SYN = {
    'propertyaddress': 'address', 'address': 'address', 'street': 'address', 'property': 'address',
    'submarket': 'city', 'city': 'city', 'neighborhood': 'city', 'location': 'city',
    'transactiondate': 'close', 'saledate': 'close', 'closedate': 'close', 'date': 'close',
    'totalsf': 'bldg', 'sf': 'bldg', 'buildingsf': 'bldg', 'size': 'bldg', 'grosssf': 'bldg', 'rsf': 'bldg',
    'saleprice': 'price', 'price': 'price', 'closeprice': 'price',
    'pricepsf': 'ppu', 'psf': 'ppu', 'priceperunit': 'ppu', 'ppu': 'ppu', 'pricesf': 'ppu',
    'caprate': 'cap', 'cap': 'cap',
    'units': 'units', 'unitcount': 'units', 'numunits': 'units',
    'assettype': 'asset_type', 'type': 'asset_type',
    'dealoverview': 'notes', 'overview': 'notes', 'notes': 'notes', 'comments': 'notes',
    'name': 'name', 'comp': 'name', 'building': 'name',
}
_FOOTER = ('subject', 'source', 'note:', 'notes:')


def _rows(sheet, limit=400):
    for r in sheet.iter_rows(max_row=limit, values_only=True):
        yield list(r)


def parse_comps(data, kind='sale'):
    wb = _wb(data)
    # pick the sheet: prefer one whose name matches the kind, else the first.
    ws = None
    for name in wb.sheetnames:
        if kind in name.lower():
            ws = wb[name]
            break
    ws = ws or wb[wb.sheetnames[0]]

    grid = _rows(ws)
    all_rows = list(grid)
    # find the header row: the first row with >=3 recognised comp headers.
    hidx, keys = None, None
    for i, row in enumerate(all_rows[:15]):
        mapped = [COMP_SYN.get(norm(c)) for c in row]
        if sum(1 for m in mapped if m) >= 3:
            hidx, keys = i, mapped
            break
    if hidx is None:
        return {'rows': []}

    out = []
    for row in all_rows[hidx + 1:]:
        first = norm(row[0]) if row and row[0] is not None else ''
        if not any(c not in (None, '') for c in row):
            break                                  # blank row ends the table
        if first.startswith(_FOOTER) or (first and 'subjectproperty' in first):
            break                                  # footer / sources line
        rec = {}
        for j, key in enumerate(keys):
            if not key or j >= len(row) or row[j] in (None, ''):
                continue
            v = row[j]
            if key == 'price':
                rec[key] = money(v)
            elif key == 'ppu':
                rec[key] = money(v) if num(v) and num(v) > 100 else money2(v)
            elif key == 'bldg':
                n = num(v)
                rec[key] = f'{round(n):,}' if n is not None else str(v).strip()
            elif key == 'cap':
                rec[key] = as_pct(v)
            elif key == 'close':
                rec[key] = date_str(v)
            else:
                rec[key] = str(v).strip()
        if rec.get('address') or rec.get('name'):
            rec.setdefault('name', rec.get('address', ''))
            out.append(rec)
    return {'rows': out}


# --------------------------------------------------------------- financials
# Deal Summary label -> console property field (normalised label match).
DEAL_FIELDS = {
    'propertyaddress': 'addr',
    'yearbuilt': 'yr',
    'totalrentablesf': 'sf',
    'netoperatingincome': 'noi',
    'totalunits': 'units',
}


def _deal_summary(ws):
    prop, extra, highlights = {}, {}, []
    in_highlights = False
    for row in _rows(ws, limit=60):
        cells = list(row) + [None] * (8 - len(row))
        # two label/value pairs: (A,B) and (E,F)
        for li, vi in ((0, 1), (4, 5)):
            label, value = cells[li], cells[vi]
            key = norm(label)
            if not key:
                continue
            if key in DEAL_FIELDS and value not in (None, ''):
                prop[DEAL_FIELDS[key]] = value
            elif key.startswith('occupancyrate') and value not in (None, ''):
                prop['occ'] = value
            elif key == 'neighborhood':
                extra['neighborhood'] = str(value).strip() if value else ''
            elif key == 'zipcode':
                extra['zip'] = str(value).strip() if value else ''
            elif key == 'zoning':
                extra['zoning'] = str(value).strip() if value else ''
            elif key == 'dealhighlights':
                in_highlights = True
        # Deal Highlights: bullet lines in col E below the header
        e = cells[4]
        if in_highlights and isinstance(e, str) and e.strip().startswith(('•', '-', '·')):
            highlights.append(e.strip().lstrip('•-·').strip())
    # city = neighborhood + zip
    city = extra.get('neighborhood', '')
    if extra.get('zip'):
        city = f'{city} {extra["zip"]}'.strip()
    if city:
        prop['city'] = city
    # normalise numbers to the console's expected shapes
    if 'sf' in prop:
        n = num(prop['sf']);  prop['sf'] = str(round(n)) if n else ''
    if 'noi' in prop:
        n = num(prop['noi']); prop['noi'] = str(round(n)) if n else ''
    if 'units' in prop:
        n = num(prop['units']); prop['units'] = str(round(n)) if n else ''
    if 'occ' in prop:
        prop['occ'] = as_pct(prop['occ']).rstrip('%')
    if 'yr' in prop:
        n = num(prop['yr']); prop['yr'] = str(round(n)) if n else str(prop['yr']).strip()
    if 'addr' in prop:
        prop['addr'] = str(prop['addr']).strip()
    prop.update({k: extra[k] for k in ('zoning',) if k in extra})
    return prop, highlights


def _operating(ws):
    """OpEx sheet -> operating statement lines with $ and $/SF, grouped.

    Stops after the Net Operating Income line so the metric block below it
    (Occupancy, Expense Ratio, WALT, …) isn't swept into the P&L."""
    lines, group = [], ''
    for row in _rows(ws, limit=60):
        cells = list(row) + [None] * (4 - len(row))
        label, amount, psf = cells[0], cells[1], cells[2]
        lab = str(label).strip() if label else ''
        if not lab:
            continue
        up = lab.upper()
        if up in ('REVENUES', 'REVENUE', 'OPERATING EXPENSES'):
            group = 'rev' if 'REVENUE' in up else 'exp'
            continue
        if num(amount) is None:
            continue
        low = lab.lower()
        is_total = low.startswith('total')
        g = group
        if 'net operating income' in low:
            g, is_total = 'noi', True
        lines.append({'label': lab, 'amount': money(amount),
                      'psf': money2(psf), 'group': g, 'total': is_total})
        if g == 'noi':
            break
    return lines


CASHFLOW_ROWS = ('total rental revenue', 'total expense recoveries',
                 'effective gross revenue', 'total operating expenses',
                 'net operating income')


def _cashflow(ws, max_years=10):
    rows = list(_rows(ws, limit=70))
    # find the "For the Year Ending" row to label year columns
    years, ystart = [], 3            # col D (0-indexed 3) = Year 1
    for row in rows:
        if row and isinstance(row[0], str) and 'year ending' in row[0].lower():
            for c in row[ystart:ystart + max_years]:
                years.append(year_of(c))
            break
    if not years:
        years = list(range(1, max_years + 1))
    out_rows = []
    for row in rows:
        lab = str(row[0]).strip() if row and row[0] else ''
        if lab.lower() in CASHFLOW_ROWS:
            vals = [money(row[ystart + k]) if ystart + k < len(row) else ''
                    for k in range(len(years))]
            out_rows.append({'label': lab, 'values': vals})
    return {'years': [str(y) for y in years if y], 'rows': out_rows}


def parse_financials(data):
    wb = _wb(data)
    by = {norm(n): n for n in wb.sheetnames}

    def sheet(*cands):
        for c in cands:
            if c in by:
                return wb[by[c]]
        return None

    prop, highlights = {}, []
    ds = sheet('dealsummary', 'summary')
    if ds is not None:
        prop, highlights = _deal_summary(ds)
    operating = []
    ox = sheet('opex', 'operatingexpenses')
    if ox is not None:
        operating = _operating(ox)
    cashflow = {'years': [], 'rows': []}
    pf = sheet('proforma')
    if pf is not None:
        cashflow = _cashflow(pf)
    return {'property': prop, 'highlights': highlights,
            'operating': operating, 'cashflow': cashflow}
